from selenium import webdriver
from selenium.webdriver.common.by import By
from openpyxl import Workbook
import time

option = webdriver.ChromeOptions()
driver = webdriver.Chrome(options=option)

driver.get("https://opensource-demo.orangehrmlive.com/web/index.php/auth/login")
time.sleep(10)
driver.find_element(By.XPATH, '//*[@id="app"]/div[1]/div/div[1]/div/div[2]/div[2]/form/div[1]/div/div[2]/input').send_keys("Admin")
time.sleep(7)
driver.find_element(By.XPATH, '/html/body/div/div[1]/div/div[1]/div/div[2]/div[2]/form/div[2]/div/div[2]/input').send_keys("admin123")
time.sleep(5)
driver.find_element(By.XPATH, '/html/body/div/div[1]/div/div[1]/div/div[2]/div[2]/form/div[3]/button').click()
time.sleep(5)
driver.find_element(By.XPATH, '/html/body/div/div[1]/div[1]/aside/nav/div[2]/ul/li[1]/a').click()
time.sleep(5)

admin_rows = driver.find_elements(By.XPATH, '/html/body/div/div[1]/div[2]/div[2]/div/div[2]/div[3]/div/div[2]/div')

wb = Workbook()
admin_ws = wb.active


admin_ws.title = "Admin Data"
admin_ws.append(["Username", "User_role", "Employee_name", "Status"])  

for i in range(1, len(admin_rows)+1):
        Username = driver.find_element(By.XPATH, '/html/body/div/div[1]/div[2]/div[2]/div/div[2]/div[3]/div/div[2]/div[' +str(i)+ ']/div/div[2]').text
        User_role = driver.find_element(By.XPATH, '/html/body/div/div[1]/div[2]/div[2]/div/div[2]/div[3]/div/div[2]/div[' +str(i)+ ']/div/div[3]').text
        Employee_name = driver.find_element(By.XPATH, '/html/body/div/div[1]/div[2]/div[2]/div/div[2]/div[3]/div/div[2]/div[' +str(i)+ ']/div/div[4]').text
        Status = driver.find_element(By.XPATH, '/html/body/div/div[1]/div[2]/div[2]/div/div[2]/div[3]/div/div[2]/div[' +str(i)+ ']/div/div[5]').text
    

        admin_ws.append([Username, User_role, Employee_name, Status])  


driver.find_element(By.XPATH, '/html/body/div/div[1]/div[1]/aside/nav/div[2]/ul/li[2]/a').click()
time.sleep(5)

pim_rows = driver.find_elements(By.XPATH, '/html/body/div/div[1]/div[2]/div[2]/div/div[2]/div[3]/div/div[2]/div')

pim_ws = wb.create_sheet(title="PIM Data")
pim_ws.append(["ID", "First_name", "Last_name", "Job_title", "Employee_status", "Subunit"])  


for i in range(1, len(pim_rows)+1):
    id = driver.find_element(By.XPATH, '/html/body/div/div[1]/div[2]/div[2]/div/div[2]/div[3]/div/div[2]/div[' +str(i)+ ']/div/div[2]').text
    First_name = driver.find_element(By.XPATH, '/html/body/div/div[1]/div[2]/div[2]/div/div[2]/div[3]/div/div[2]/div[' +str(i)+ ']/div/div[3]').text
    Last_name = driver.find_element(By.XPATH, '/html/body/div/div[1]/div[2]/div[2]/div/div[2]/div[3]/div/div[2]/div[' +str(i)+ ']/div/div[4]').text
    Job_title = driver.find_element(By.XPATH, '/html/body/div/div[1]/div[2]/div[2]/div/div[2]/div[3]/div/div[2]/div[' +str(i)+ ']/div/div[5]').text
    Employment_status = driver.find_element(By.XPATH, '/html/body/div/div[1]/div[2]/div[2]/div/div[2]/div[3]/div/div[2]/div[' +str(i)+ ']/div/div[6]').text
    Subunit = driver.find_element(By.XPATH, '/html/body/div/div[1]/div[2]/div[2]/div/div[2]/div[3]/div/div[2]/div[' +str(i)+ ']/div/div[7]').text

    pim_ws.append([id, First_name, Last_name, Job_title, Employment_status, Subunit])  


driver.find_element(By.XPATH, '/html/body/div/div[1]/div[1]/aside/nav/div[2]/ul/li[3]/a ').click()
time.sleep(5)

leave_rows = driver.find_elements(By.XPATH, '/html/body/div/div[1]/div[2]/div[2]/div/div[2]/div[2]/div/div[2]/div')

Leave_ws = wb.create_sheet(title="Leave Data")
Leave_ws.append(["Date", "Employee_name", "Leave_type", "Leave_balance", "Number_of_days", "status", "Comments"])  


for i in range(1, len(leave_rows)+1):
    Date= driver.find_element(By.XPATH, '/html/body/div/div[1]/div[2]/div[2]/div/div[2]/div[2]/div/div[2]/div[' +str(i)+ ']/div/div[2]').text
    Employee_name = driver.find_element(By.XPATH, '/html/body/div/div[1]/div[2]/div[2]/div/div[2]/div[2]/div/div[2]/div[' +str(i)+ ']/div/div[3]').text
    Leave_type = driver.find_element(By.XPATH, '/html/body/div/div[1]/div[2]/div[2]/div/div[2]/div[2]/div/div[2]/div[' +str(i)+ ']/div/div[4]').text
    Leave_balance = driver.find_element(By.XPATH, '/html/body/div/div[1]/div[2]/div[2]/div/div[2]/div[2]/div/div[2]/div[' +str(i)+ ']/div/div[5]').text
    Number_of_days = driver.find_element(By.XPATH, '/html/body/div/div[1]/div[2]/div[2]/div/div[2]/div[2]/div/div[2]/div[' +str(i)+ ']/div/div[6]').text
    status = driver.find_element(By.XPATH, '/html/body/div/div[1]/div[2]/div[2]/div/div[2]/div[2]/div/div[2]/div[' +str(i)+ ']/div/div[7]').text
    Comments = driver.find_element(By.XPATH, '/html/body/div/div[1]/div[2]/div[2]/div/div[2]/div[2]/div/div[2]/div[' +str(i)+ ']/div/div[8]').text

    Leave_ws.append([Date, Employee_name, Leave_type, Leave_balance, Number_of_days, status, Comments])  

driver.find_element(By.XPATH, '/html/body/div/div[1]/div[1]/aside/nav/div[2]/ul/li[4]/a').click()
time.sleep(5)

time_rows = driver.find_elements(By.XPATH, '/html/body/div/div[1]/div[2]/div[2]/div/div[2]/div[3]/div/div[2]/div')

time_ws = wb.create_sheet(title="Time Data")
time_ws.append(["Employee_name", "Timesheet_period"])  


for i in range(1, len(time_rows)+1):
    Employee_name = driver.find_element(By.XPATH, '/html/body/div/div[1]/div[2]/div[2]/div/div[2]/div[3]/div/div[2]/div[' +str(i)+ ']/div/div[1]').text
    Timesheet = driver.find_element(By.XPATH, '/html/body/div/div[1]/div[2]/div[2]/div/div[2]/div[3]/div/div[2]/div[' +str(i)+ ']/div/div[2]').text

    
    time_ws.append([Employee_name, Timesheet])  

driver.find_element(By.XPATH, '/html/body/div/div[1]/div[1]/aside/nav/div[2]/ul/li[2]/a').click()
time.sleep(5)

recruitment_rows = driver.find_elements(By.XPATH, '/html/body/div/div[1]/div[2]/div[2]/div/div[2]/div[3]/div/div[2]/div')

recruitment_ws = wb.create_sheet(title="Recruitment Data")
recruitment_ws.append(["Vacancy", "Candidate", "Hiring Manager", "Date of Application", "Status"])  


for i in range(1, len(recruitment_rows)+1):
    Vacancy = driver.find_element(By.XPATH, '/html/body/div/div[1]/div[2]/div[2]/div/div[2]/div[3]/div/div[2]/div[' +str(i)+ ']/div/div[2]').text
    Candidate = driver.find_element(By.XPATH, '/html/body/div/div[1]/div[2]/div[2]/div/div[2]/div[3]/div/div[2]/div[' +str(i)+ ']/div/div[3]').text
    Hm = driver.find_element(By.XPATH, '/html/body/div/div[1]/div[2]/div[2]/div/div[2]/div[3]/div/div[2]/div[' +str(i)+ ']/div/div[4]').text
    Doa = driver.find_element(By.XPATH, '/html/body/div/div[1]/div[2]/div[2]/div/div[2]/div[3]/div/div[2]/div[' +str(i)+ ']/div/div[5]').text
    Status = driver.find_element(By.XPATH, '/html/body/div/div[1]/div[2]/div[2]/div/div[2]/div[3]/div/div[2]/div[' +str(i)+ ']/div/div[6]').text

    recruitment_ws.append([Vacancy, Candidate, Hm, Doa, Status])  


wb.save("admin_data.xlsx")





